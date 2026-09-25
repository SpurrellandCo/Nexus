#!/usr/bin/env python3
"""Make branded PowerPoint decks with this kit.

  make_pptx.py outline.md   -o deck.pptx      build a deck from a slides outline
  make_pptx.py existing.pptx -o branded.pptx  re-brand an existing deck: keeps each slide's
                                              title, text, tables (as text), and speaker notes;
                                              reports pictures/charts it could not carry over

Outline format: slides separated by a line containing only '---'. In each slide:
'# Title' (or '## Title'), an optional subtitle line, '- bullets' (indent 2 spaces for
sub-points), and 'Notes: ...' for speaker notes. A first slide without bullets becomes
the title slide.

If the user supplied their own slide template (brand.json templates.pptx source "user"),
its masters and layouts are used; otherwise the brand look is drawn from brand.json.
"""
import argparse
import sys
from pathlib import Path

from brandkit import color, kit_file, load_brand, parse_slides, require, rgb, template

require("pptx", "python-pptx")
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

WIDTH, HEIGHT = Inches(13.333), Inches(7.5)


def _rgb(brand, name, fallback="text"):
    return RGBColor(*rgb(color(brand, name, fallback)))


def _fonts(brand):
    return brand["fonts"]["heading"], brand["fonts"].get("body", brand["fonts"]["heading"])


def _textbox(slide, box, lines, font, size, rgb_color, bold=False, bullets=False):
    frame = slide.shapes.add_textbox(*box).text_frame
    frame.word_wrap = True
    for i, (level, text) in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = ("    " * level + "• " if bullets else "") + text
        run.font.name, run.font.size, run.font.bold = font, Pt(size - 2 * level), bold
        run.font.color.rgb = rgb_color
        paragraph.space_after = Pt(8)
    return frame


def _logo(slide, brand, left, top, height):
    logo = kit_file(brand.get("logo"))
    if logo:
        slide.shapes.add_picture(str(logo), left, top, height=height)


def _blank_layout(prs):
    named = [layout for layout in prs.slide_layouts if layout.name.lower() == "blank"]
    return named[0] if named else prs.slide_layouts[len(prs.slide_layouts) - 1]


def remove_all_slides(prs):
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def new_deck(brand, use_user_template=True):
    path, source = template(brand, "pptx")
    if use_user_template and path and source == "user":
        prs = Presentation(str(path))
        remove_all_slides(prs)
        return prs, "template"
    prs = Presentation()
    prs.slide_width, prs.slide_height = WIDTH, HEIGHT
    return prs, "brand"


def add_title_slide(prs, mode, brand, title, subtitle):
    if mode == "template":
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        if slide.shapes.title is not None:
            slide.shapes.title.text = title
        others = [p for p in slide.placeholders if p.placeholder_format.idx != 0 and p.has_text_frame]
        if subtitle and others:
            others[0].text_frame.text = subtitle
        return slide
    heading_font, body_font = _fonts(brand)
    slide = prs.slides.add_slide(_blank_layout(prs))
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    band.fill.solid()
    band.fill.fore_color.rgb = _rgb(brand, "primary")
    band.line.fill.background()
    white = _rgb(brand, "background", "text")
    frame = _textbox(slide, (Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.6)), [(0, title)], heading_font, 44, white, bold=True)
    frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    if subtitle:
        _textbox(slide, (Inches(0.9), Inches(4.1), Inches(11.5), Inches(1.0)), [(0, subtitle)], body_font, 22, white)
    _logo(slide, brand, prs.slide_width - Inches(2.0), prs.slide_height - Inches(1.3), Inches(0.8))
    return slide


def add_content_slide(prs, mode, brand, title, bullets):
    heading_font, body_font = _fonts(brand)
    if mode == "template":
        layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        if slide.shapes.title is not None:
            slide.shapes.title.text = title
        body = [p for p in slide.placeholders if p.placeholder_format.idx != 0 and p.has_text_frame]
        if body:
            frame = body[0].text_frame
            for i, (level, text) in enumerate(bullets):
                paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                paragraph.text, paragraph.level = text, min(level, 4)
            return slide
        _textbox(slide, (Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.2)), bullets, body_font, 20, _rgb(brand, "text"), bullets=True)
        return slide
    slide = prs.slides.add_slide(_blank_layout(prs))
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.25), prs.slide_height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = _rgb(brand, "accent", "primary")
    bar.line.fill.background()
    _textbox(slide, (Inches(0.8), Inches(0.5), Inches(10.5), Inches(1.2)), [(0, title)], heading_font, 32, _rgb(brand, "primary"), bold=True)
    if bullets:
        _textbox(slide, (Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.2)), bullets, body_font, 20, _rgb(brand, "text"), bullets=True)
    _logo(slide, brand, prs.slide_width - Inches(1.6), Inches(0.45), Inches(0.6))
    return slide


def add_slides(prs, mode, brand, slides):
    for index, spec in enumerate(slides):
        if index == 0 and not spec["bullets"]:
            slide = add_title_slide(prs, mode, brand, spec["title"], spec["subtitle"])
        else:
            bullets = ([(0, spec["subtitle"])] if spec["subtitle"] else []) + spec["bullets"]
            slide = add_content_slide(prs, mode, brand, spec["title"], bullets)
        if spec["notes"]:
            slide.notes_slide.notes_text_frame.text = spec["notes"]


def _shape_text(shape, lines):
    if shape.has_text_frame:
        lines.extend((p.level, p.text.strip()) for p in shape.text_frame.paragraphs if p.text.strip())
    elif getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            lines.append((0, " | ".join(cell.text.strip() for cell in row.cells)))
    elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for inner in shape.shapes:
            _shape_text(inner, lines)


def extract(source):
    """Slide specs from an existing deck, plus what couldn't be carried over."""
    slides, skipped = [], {"pictures": [], "charts": []}
    for number, slide in enumerate(Presentation(str(source)).slides, start=1):
        title_shape = slide.shapes.title
        spec = {"title": title_shape.text.strip() if title_shape is not None else "", "subtitle": "", "bullets": [], "notes": ""}
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or getattr(shape, "image", None) is not None and shape.is_placeholder:
                skipped["pictures"].append(number)
            elif getattr(shape, "has_chart", False) and shape.has_chart:
                skipped["charts"].append(number)
            else:
                _shape_text(shape, spec["bullets"])
        if slide.has_notes_slide:
            spec["notes"] = slide.notes_slide.notes_text_frame.text.strip()
        slides.append(spec)
    return slides, skipped


def build(source, output, brand):
    prs, mode = new_deck(brand)
    slides = parse_slides(Path(source).read_text())
    add_slides(prs, mode, brand, slides)
    prs.save(str(output))
    return len(slides)


def rebrand(source, output, brand):
    slides, skipped = extract(source)
    prs, mode = new_deck(brand)
    add_slides(prs, mode, brand, slides)
    prs.save(str(output))
    return len(slides), skipped


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("input", help="slides outline (.md/.txt) or an existing .pptx to re-brand")
    parser.add_argument("-o", "--output", help="output .pptx (default: <input>-branded.pptx)")
    args = parser.parse_args(argv)
    source = Path(args.input)
    if not source.exists():
        sys.exit(f"Not found: {source}")
    output = Path(args.output) if args.output else source.with_name(f"{source.stem}-branded.pptx")
    brand = load_brand()
    if source.suffix.lower() in (".pptx", ".potx"):
        count, skipped = rebrand(source, output, brand)
        print(f"Re-branded {source.name} -> {output} ({count} slides; titles, text, tables, and notes kept)")
        for kind in ("pictures", "charts"):
            if skipped[kind]:
                slides = ", ".join(str(n) for n in sorted(set(skipped[kind])))
                print(f"Not carried over: {kind} on slide(s) {slides}. Re-add them from the original deck.")
    else:
        print(f"Wrote {output} ({build(source, output, brand)} slides)")


if __name__ == "__main__":
    main()
