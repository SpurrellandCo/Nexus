#!/usr/bin/env python3
"""Make branded Word documents with this kit.

  make_docx.py content.md  -o report.docx    build a new document from markdown (or .txt)
  make_docx.py existing.docx -o branded.docx re-brand an existing document: keeps every word,
                                             table, and image; applies the brand's fonts,
                                             colors, and logo
  make_docx.py existing.pdf -o branded.docx  rebuild a PDF's text (headings, paragraphs, bullets)
                                             as a branded document; layout and images are not
                                             carried over (reported)

The base is templates/word-template.docx when present (a user's own letterhead
keeps its headers and footers; its sample body text is dropped).
"""
import argparse
import sys
from pathlib import Path

from brandkit import color, kit_file, load_brand, parse_markdown, inline_runs, require, rgb, template

docx = require("docx", "python-docx")
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

HEADING_STYLES = {"Title": (28, "primary"), "Heading 1": (20, "primary"),
                  "Heading 2": (15, "secondary"), "Heading 3": (12.5, "text")}


def _style(doc, name):
    try:
        return doc.styles[name]
    except KeyError:
        return None


def _set_font(font_owner, name, size=None, hex_color=None):
    font_owner.font.name = name
    rpr = font_owner.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)
    if size:
        font_owner.font.size = Pt(size)
    if hex_color:
        font_owner.font.color.rgb = RGBColor(*rgb(hex_color))


LIST_STYLES = ("List Bullet", "List Number")


def _recolor_title_rule(doc, brand):
    """Word's default Title has a blue rule under it; make it the brand accent."""
    title = _style(doc, "Title")
    borders = title.element.pPr.find(qn("w:pBdr")) if title is not None and title.element.pPr is not None else None
    if borders is not None:
        for edge in borders:
            edge.set(qn("w:color"), color(brand, "accent", "primary").lstrip("#"))


def _tidy_lists(doc):
    """Explicit indents and tight spacing, so lists look right in every viewer."""
    for base in LIST_STYLES:
        for level, name in enumerate((base, f"{base} 2", f"{base} 3")):
            style = _style(doc, name)
            if style is None:
                continue
            fmt = style.paragraph_format
            fmt.left_indent, fmt.first_line_indent = Cm(0.63 * (level + 1)), Cm(-0.63)
            fmt.space_before, fmt.space_after = Pt(0), Pt(3)


def apply_brand_styles(doc, brand):
    heading_font, body_font = brand["fonts"]["heading"], brand["fonts"].get("body", brand["fonts"]["heading"])
    normal = _style(doc, "Normal")
    if normal is not None:
        _set_font(normal, body_font, 11, color(brand, "text"))
    for name, (size, color_name) in HEADING_STYLES.items():
        style = _style(doc, name)
        if style is not None:
            _set_font(style, heading_font, size, color(brand, color_name))
    _recolor_title_rule(doc, brand)
    _tidy_lists(doc)


def add_logo_header(doc, brand):
    logo = kit_file(brand.get("logo"))
    if not logo:
        return
    header = doc.sections[0].header
    if "graphic" in header._element.xml:
        return  # already has a logo or image
    paragraph = header.add_paragraph() if header.paragraphs[0].text.strip() else header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run().add_picture(str(logo), height=Cm(1.2))


def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _add_runs(paragraph, text):
    for chunk, bold, italic in inline_runs(text):
        run = paragraph.add_run(chunk)
        run.bold, run.italic = bold or None, italic or None


def _shade(cell, hex_color):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), hex_color.lstrip("#"))
    cell._tc.get_or_add_tcPr().append(shading)


def _list_paragraph(doc, style_name, level, text):
    style = style_name if level == 0 else f"{style_name} {min(level + 1, 3)}"
    paragraph = doc.add_paragraph(style=style if _style(doc, style) else None)
    if not _style(doc, style):
        paragraph.add_run("  " * level + ("• " if "Bullet" in style_name else ""))
    _add_runs(paragraph, text)


def _table_borders(table, hex_color):
    """Light horizontal rules in a brand color; no vertical lines."""
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        visible = edge in ("top", "bottom", "insideH")
        element.set(qn("w:val"), "single" if visible else "nil")
        if visible:
            element.set(qn("w:sz"), "4")
            element.set(qn("w:color"), hex_color.lstrip("#"))
        borders.append(element)
    table._tbl.tblPr.append(borders)


def _full_width(table, total_emu):
    """Explicit table width and fixed layout (python-docx leaves both on 'auto')."""
    props = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        existing = props.find(qn(tag))
        if existing is not None:
            props.remove(existing)
    table_width = OxmlElement("w:tblW")
    table_width.set(qn("w:w"), str(int(total_emu / 635)))  # EMU -> twentieths of a point
    table_width.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    props.append(table_width)
    props.append(layout)


def _add_table(doc, brand, rows):
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.autofit = False
    section = doc.sections[-1]
    column = int((section.page_width - section.left_margin - section.right_margin) / width)
    _table_borders(table, color(brand, "muted", "text"))
    for col in table.columns:  # grid widths too: some viewers ignore per-cell widths
        col.width = column
    _full_width(table, column * width)
    for r, row in enumerate(rows):
        for c in range(width):
            cell = table.cell(r, c)
            cell.width = column
            cell.text = ""
            cell.paragraphs[0].paragraph_format.space_before = Pt(3)
            cell.paragraphs[0].paragraph_format.space_after = Pt(3)
            _add_runs(cell.paragraphs[0], row[c] if c < len(row) else "")
            if r == 0:
                _shade(cell, color(brand, "primary"))
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(*rgb(color(brand, "background", "text")))


def _heading_style(doc, blocks, index, level):
    """A single leading '# ' becomes the Title; the rest shift up so '## ' is Heading 1."""
    level_ones = [b for b in blocks if b["type"] == "heading" and b["level"] == 1]
    title_first = blocks[0]["type"] == "heading" and blocks[0]["level"] == 1 and len(level_ones) == 1
    if title_first and index == 0 and _style(doc, "Title"):
        return "Title"
    return f"Heading {max(1, min(level - (1 if title_first else 0), 3))}"


def render_blocks(doc, brand, blocks):
    for index, block in enumerate(blocks):
        kind = block["type"]
        if kind == "heading":
            paragraph = doc.add_paragraph(style=_heading_style(doc, blocks, index, block["level"]))
            _add_runs(paragraph, block["text"])
        elif kind == "para":
            _add_runs(doc.add_paragraph(), block["text"])
        elif kind in ("bullets", "numbers"):
            for level, text in block["items"]:
                _list_paragraph(doc, "List Bullet" if kind == "bullets" else "List Number", level, text)
        elif kind == "table":
            _add_table(doc, brand, block["rows"])
        elif kind == "quote":
            paragraph = doc.add_paragraph(style="Quote" if _style(doc, "Quote") else None)
            paragraph.paragraph_format.space_before = Pt(12)
            _add_runs(paragraph, block["text"])
            if not _style(doc, "Quote"):
                for run in paragraph.runs:
                    run.italic = True


def build_blocks(blocks, output, brand):
    base, _source = template(brand, "docx")
    doc = docx.Document(str(base)) if base else docx.Document()
    if base:
        clear_body(doc)
    apply_brand_styles(doc, brand)
    add_logo_header(doc, brand)
    render_blocks(doc, brand, blocks)
    doc.save(str(output))


def build(source, output, brand):
    build_blocks(parse_markdown(Path(source).read_text()), output, brand)


def pdf_to_blocks(source):
    """Blocks from a PDF's text, or exit with a clear message when it has none."""
    from pdf_text import pdf_blocks
    blocks, report = pdf_blocks(source)
    if not blocks:
        sys.exit(f"No text found in {Path(source).name} (it may be a scanned image). "
                 "OCR isn't supported; run the PDF through an OCR tool first, then re-brand the result.")
    return blocks, report


def _clear_direct_fonts(paragraph, keep_color=False):
    for run in paragraph.runs:
        run.font.name = None
        rpr = run._element.rPr
        if rpr is not None and rpr.find(qn("w:rFonts")) is not None:
            rpr.remove(rpr.find(qn("w:rFonts")))
        if not keep_color:
            run.font.color.rgb = None


def rebrand(source, output, brand):
    doc = docx.Document(str(source))
    apply_brand_styles(doc, brand)
    paragraphs = list(doc.paragraphs) + [p for t in doc.tables for row in t.rows for cell in row.cells for p in cell.paragraphs]
    for paragraph in paragraphs:
        is_heading = paragraph.style is not None and (paragraph.style.name in HEADING_STYLES or paragraph.style.name.startswith("Heading"))
        _clear_direct_fonts(paragraph, keep_color=not is_heading)
    add_logo_header(doc, brand)
    doc.save(str(output))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("input", help="markdown/.txt content, or an existing .docx to re-brand")
    parser.add_argument("-o", "--output", help="output .docx (default: <input>-branded.docx)")
    args = parser.parse_args(argv)
    source = Path(args.input)
    if not source.exists():
        sys.exit(f"Not found: {source}")
    output = Path(args.output) if args.output else source.with_name(f"{source.stem}-branded.docx")
    brand = load_brand()
    if source.suffix.lower() in (".docx", ".dotx"):
        rebrand(source, output, brand)
        print(f"Re-branded {source.name} -> {output} (content kept; brand fonts, colors, and logo applied)")
    elif source.suffix.lower() == ".pdf":
        blocks, report = pdf_to_blocks(source)
        build_blocks(blocks, output, brand)
        print(f"Rebuilt {source.name} ({report['pages']} page(s)) as {output}: text, headings, and lists kept; "
              f"original layout not kept" + (f"; {report['images']} image(s) to re-add by hand" if report["images"] else ""))
    else:
        build(source, output, brand)
        print(f"Wrote {output}")


if __name__ == "__main__":
    main()
