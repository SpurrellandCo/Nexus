#!/usr/bin/env python3
"""Make branded Excel spreadsheets with this kit.

  make_xlsx.py data.csv -o table.xlsx          a branded spreadsheet from CSV (or the first
                                               markdown table in a .md file)
  make_xlsx.py existing.xlsx -o branded.xlsx   re-brand a workbook: keeps every value, formula,
                                               and number format; applies the brand fonts,
                                               header colors, and sheet tab colors

Charts, images, and pivot tables can't be carried through (the Excel library used here drops
them), so re-branding lists any it finds; the original file is never changed.
If the user supplied their own spreadsheet template (brand.json templates.xlsx source "user"),
new spreadsheets are built on it.
"""
import argparse
import csv
import re
import sys
import zipfile
from copy import copy
from pathlib import Path

from brandkit import color, load_brand, parse_markdown, plain, require, rgb, template

require("openpyxl", "openpyxl")
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

LOST_PARTS = {"charts": "xl/charts/", "images": "xl/media/", "pivot tables": "xl/pivotTables/", "macros": "xl/vbaProject.bin"}


def _hex(brand, name, fallback="text"):
    return color(brand, name, fallback).lstrip("#").upper()


def _tint(hex_color, amount=0.9):
    r, g, b = rgb(hex_color)
    return "".join(f"{round(v + (255 - v) * amount):02X}" for v in (r, g, b))


def _fonts(brand):
    return brand["fonts"]["heading"], brand["fonts"].get("body", brand["fonts"]["heading"])


def _number(value):
    text = value.strip()
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    if re.fullmatch(r"-?\d*\.\d+", text):
        return float(text)
    return value


def style_header_row(ws, row, brand):
    heading_font, _ = _fonts(brand)
    fill = PatternFill("solid", fgColor=_hex(brand, "primary"))
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.fill = fill
        cell.font = Font(name=heading_font, bold=True, color=_hex(brand, "background", "text"), size=cell.font.size or 11)
        cell.alignment = Alignment(vertical="center")


def write_table(ws, rows, brand):
    heading_font, body_font = _fonts(brand)
    band = PatternFill("solid", fgColor=_tint(color(brand, "primary")))
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=_number(value) if r > 1 else value)
            if r > 1:
                cell.font = Font(name=body_font, color=_hex(brand, "text"))
                if r % 2 == 1:
                    cell.fill = band
    style_header_row(ws, 1, brand)
    for c in range(1, max(len(row) for row in rows) + 1):
        longest = max(len(str(row[c - 1])) if c - 1 < len(row) else 0 for row in rows)
        ws.column_dimensions[get_column_letter(c)].width = min(60, max(10, longest * 1.2 + 2))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_properties.tabColor = _hex(brand, "primary")


def read_rows(source):
    if source.suffix.lower() == ".csv":
        with source.open(newline="", encoding="utf-8-sig") as handle:
            return [row for row in csv.reader(handle) if any(cell.strip() for cell in row)]
    tables = [b for b in parse_markdown(source.read_text()) if b["type"] == "table"]
    if not tables:
        sys.exit(f"No table found in {source.name}. Give a CSV file, or markdown with a | table |.")
    return [[plain(cell) for cell in row] for row in tables[0]["rows"]]


def build(source, output, brand):
    rows = read_rows(Path(source))
    path, origin = template(brand, "xlsx")
    if path and origin == "user":
        wb = load_workbook(str(path))
        ws = wb.worksheets[0]
        if ws.max_row:
            ws.delete_rows(1, ws.max_row)
        wb.active = 0
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = brand["name"][:31]
    write_table(ws, rows, brand)
    wb.save(str(output))
    return len(rows) - 1


def lost_parts(source):
    with zipfile.ZipFile(source) as archive:
        names = archive.namelist()
    return [label for label, prefix in LOST_PARTS.items() if any(n.startswith(prefix) for n in names)]


def rebrand(source, output, brand):
    heading_font, body_font = _fonts(brand)
    wb = load_workbook(str(source))
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = _hex(brand, "primary")
        if ws.max_row < 1 or ws.max_column < 1:
            continue
        header = ws.min_row
        for row in ws.iter_rows(min_row=header + 1):
            for cell in row:
                if cell.value is not None:
                    font = copy(cell.font)
                    font.name = body_font
                    cell.font = font
        style_header_row(ws, header, brand)
        if ws.freeze_panes is None:
            ws.freeze_panes = f"A{header + 1}"
    wb.save(str(output))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("input", help="a .csv, a .md with a table, or an existing .xlsx to re-brand")
    parser.add_argument("-o", "--output", help="output .xlsx (default: <input>-branded.xlsx)")
    args = parser.parse_args(argv)
    source = Path(args.input)
    if not source.exists():
        sys.exit(f"Not found: {source}")
    output = Path(args.output) if args.output else source.with_name(f"{source.stem}-branded.xlsx")
    brand = load_brand()
    if source.suffix.lower() in (".xlsx", ".xlsm", ".xltx"):
        lost = lost_parts(source)
        rebrand(source, output, brand)
        print(f"Re-branded {source.name} -> {output} (values, formulas, and number formats kept)")
        if lost:
            print(f"Not carried over: {', '.join(lost)}. Re-add them from the original workbook ({source.name} is unchanged).")
    else:
        print(f"Wrote {output} ({build(source, output, brand)} data rows)")


if __name__ == "__main__":
    main()
