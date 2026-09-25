"""Tests for the brand-kit scripts every generated brand folder carries.

Run with: python3 -m unittest discover -s ~/.nexus/skills/brandcreator/tests
(also run by ~/.nexus/scripts/tests/brandcreator-kit.test.js)
"""
import json
import shutil
import struct
import subprocess
import sys
import tempfile
import unittest
import zlib
from pathlib import Path

from docx import Document
from pptx import Presentation

KIT_SCRIPTS = Path(__file__).resolve().parent.parent / "kit" / "scripts"
PRIMARY = "#1F4FD8"


def tiny_png(path, rgb=(31, 79, 216)):
    """Write a valid 2x2 PNG without extra dependencies."""
    raw = b"".join(b"\x00" + bytes(rgb) * 2 for _ in range(2))
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    png = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
    png += chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")
    path.write_bytes(png)


def make_kit(root):
    kit = Path(root) / "brandcreator-acme"
    shutil.copytree(KIT_SCRIPTS, kit / "scripts")
    (kit / "assets").mkdir()
    (kit / "templates").mkdir()
    tiny_png(kit / "assets" / "logo-primary.png")
    (kit / "assets" / "design-tokens.json").write_text("{}")
    (kit / "assets" / "design-tokens.css").write_text(":root { --color-primary: #1F4FD8; }")
    (kit / "templates" / "report-template.html").write_text('<link rel="stylesheet" href="../assets/design-tokens.css">')
    (kit / "brand-guidelines.md").write_text("# Acme brand guidelines\n")
    (kit / "SKILL.md").write_text("---\nname: brandcreator-acme\ndescription: Acme brand kit.\n---\nSee [guidelines](brand-guidelines.md).\n")
    (kit / "README.md").write_text("# Acme brand kit\nOpen [the Word template](templates/word-template.docx).\n")
    brand = {
        "name": "Acme",
        "slug": "brandcreator-acme",
        "colors": {"primary": PRIMARY, "secondary": "#0F172A", "accent": "#F59E0B",
                   "text": "#111827", "muted": "#6B7280", "background": "#FFFFFF"},
        "fonts": {"heading": "Space Grotesk", "body": "Inter"},
        "logo": "assets/logo-primary.png",
        "templates": {"html": "templates/report-template.html"},
    }
    (kit / "brand.json").write_text(json.dumps(brand, indent=2))
    return kit


def run(kit, script, *args, cwd=None):
    return subprocess.run([sys.executable, str(kit / "scripts" / script), *map(str, args)],
                          cwd=cwd or kit, capture_output=True, text=True)


def rgb_hex(color_format):
    return f"#{color_format.rgb}" if color_format.rgb is not None else None


class BrandKitTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="brandkit-"))
        self.kit = make_kit(self.tmp)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def templates(self):
        r = run(self.kit, "make_templates.py")
        self.assertEqual(r.returncode, 0, r.stderr)

    # ---------- templates ----------
    def test_make_templates_creates_branded_office_templates_and_records_them(self):
        self.templates()
        brand = json.loads((self.kit / "brand.json").read_text())
        self.assertEqual(brand["templates"]["docx"], {"path": "templates/word-template.docx", "source": "generated"})
        self.assertEqual(brand["templates"]["pptx"], {"path": "templates/slides-template.pptx", "source": "generated"})
        doc = Document(self.kit / "templates" / "word-template.docx")
        self.assertEqual(doc.styles["Heading 1"].font.name, "Space Grotesk")
        self.assertEqual(rgb_hex(doc.styles["Heading 1"].font.color), PRIMARY)
        self.assertEqual(doc.styles["Normal"].font.name, "Inter")
        self.assertIn("graphic", doc.sections[0].header._element.xml, "logo in the header")
        self.assertGreaterEqual(len(Presentation(self.kit / "templates" / "slides-template.pptx").slides), 2)

    def test_make_templates_never_overwrites_a_users_own_template(self):
        self.templates()
        own = Document()
        own.add_paragraph("MY LETTERHEAD")
        own.save(self.tmp / "own.docx")
        self.assertEqual(run(self.kit, "import_template.py", self.tmp / "own.docx").returncode, 0)
        run(self.kit, "make_templates.py")
        texts = [p.text for p in Document(self.kit / "templates" / "word-template.docx").paragraphs]
        self.assertIn("MY LETTERHEAD", texts)

    # ---------- Word ----------
    def test_markdown_to_branded_docx(self):
        self.templates()
        (self.tmp / "report.md").write_text(
            "# Q3 Report\n\nRevenue grew **12%** this quarter.\n\n## Highlights\n\n- New office\n- Two hires\n\n"
            "| Region | Growth |\n|---|---|\n| EU | 10% |\n| US | 14% |\n")
        r = run(self.kit, "make_docx.py", self.tmp / "report.md", "-o", self.tmp / "report.docx")
        self.assertEqual(r.returncode, 0, r.stderr)
        doc = Document(self.tmp / "report.docx")
        texts = [p.text for p in doc.paragraphs]
        self.assertIn("Q3 Report", texts)
        self.assertIn("Revenue grew 12% this quarter.", texts)
        self.assertTrue(any(p.style.name == "Heading 1" and p.text == "Highlights" for p in doc.paragraphs))
        self.assertTrue(any(p.text == "New office" for p in doc.paragraphs))
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(doc.tables[0].cell(2, 1).text, "14%")
        self.assertIn("graphic", doc.sections[0].header._element.xml)

    def test_rebrand_existing_docx_keeps_content_and_applies_brand(self):
        src = Document()
        h = src.add_heading("Old Title", level=1)
        h.runs[0].font.name = "Comic Sans MS"
        src.add_paragraph("Keep this exact sentence.")
        table = src.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "cell A"
        src.save(self.tmp / "old.docx")
        r = run(self.kit, "make_docx.py", self.tmp / "old.docx", "-o", self.tmp / "new.docx")
        self.assertEqual(r.returncode, 0, r.stderr)
        doc = Document(self.tmp / "new.docx")
        self.assertIn("Keep this exact sentence.", [p.text for p in doc.paragraphs])
        self.assertEqual(doc.tables[0].cell(0, 0).text, "cell A")
        heading = next(p for p in doc.paragraphs if p.text == "Old Title")
        self.assertIsNone(heading.runs[0].font.name, "direct font override removed so the brand style applies")
        self.assertEqual(doc.styles["Heading 1"].font.name, "Space Grotesk")
        self.assertIn("graphic", doc.sections[0].header._element.xml)

    def test_users_word_template_is_the_base_and_its_sample_text_is_dropped(self):
        own = Document()
        own.sections[0].header.paragraphs[0].text = "ACME LTD - 1 Main St"
        own.add_paragraph("Sample body text from the template")
        own.save(self.tmp / "letterhead.dotx.docx")
        self.assertEqual(run(self.kit, "import_template.py", self.tmp / "letterhead.dotx.docx", "--kind", "docx").returncode, 0)
        (self.tmp / "memo.md").write_text("# Memo\n\nHello.\n")
        run(self.kit, "make_docx.py", self.tmp / "memo.md", "-o", self.tmp / "memo.docx")
        doc = Document(self.tmp / "memo.docx")
        self.assertIn("ACME LTD - 1 Main St", doc.sections[0].header.paragraphs[0].text)
        self.assertNotIn("Sample body text from the template", [p.text for p in doc.paragraphs])
        self.assertIn("Hello.", [p.text for p in doc.paragraphs])

    def test_dotx_and_potx_templates_are_imported(self):
        Document().save(self.tmp / "t.docx")
        shutil.copy(self.tmp / "t.docx", self.tmp / "t.dotx")
        r = run(self.kit, "import_template.py", self.tmp / "t.dotx")
        self.assertEqual(r.returncode, 0, r.stderr)
        Document(self.kit / "templates" / "word-template.docx")  # opens as a normal document
        Presentation().save(self.tmp / "t.pptx")
        shutil.copy(self.tmp / "t.pptx", self.tmp / "t.potx")
        r = run(self.kit, "import_template.py", self.tmp / "t.potx")
        self.assertEqual(r.returncode, 0, r.stderr)
        Presentation(self.kit / "templates" / "slides-template.pptx")
        brand = json.loads((self.kit / "brand.json").read_text())
        self.assertEqual(brand["templates"]["pptx"]["source"], "user")

    # ---------- PowerPoint ----------
    def test_outline_to_branded_pptx(self):
        (self.tmp / "deck.md").write_text(
            "# Acme Launch\nSeptember 2026\n\n---\n\n## Why now\n- Market is ready\n  - EU first\n- Team is hired\n\n"
            "Notes: Mention the waitlist.\n\n---\n\n## Next steps\n- Ship beta\n")
        r = run(self.kit, "make_pptx.py", self.tmp / "deck.md", "-o", self.tmp / "deck.pptx")
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(self.tmp / "deck.pptx")
        self.assertEqual(len(prs.slides), 3)
        all_text = [s.shapes and "\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
        self.assertIn("Acme Launch", all_text[0])
        self.assertIn("Why now", all_text[1])
        self.assertIn("EU first", all_text[1])
        self.assertEqual(prs.slides[1].notes_slide.notes_text_frame.text, "Mention the waitlist.")
        self.assertTrue(any(sh.shape_type == 13 for sh in prs.slides[1].shapes), "logo picture on content slides")

    def test_rebrand_existing_pptx_keeps_text_and_notes_and_reports_pictures(self):
        src = Presentation()
        s1 = src.slides.add_slide(src.slide_layouts[1])
        s1.shapes.title.text = "Old deck title"
        s1.placeholders[1].text_frame.text = "Point one"
        s1.notes_slide.notes_text_frame.text = "Old notes"
        s2 = src.slides.add_slide(src.slide_layouts[5])
        s2.shapes.title.text = "Chart page"
        tiny_png(self.tmp / "pic.png")
        s2.shapes.add_picture(str(self.tmp / "pic.png"), 0, 0)
        src.save(self.tmp / "old.pptx")
        r = run(self.kit, "make_pptx.py", self.tmp / "old.pptx", "-o", self.tmp / "new.pptx")
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(self.tmp / "new.pptx")
        text = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
        self.assertIn("Old deck title", text)
        self.assertIn("Point one", text)
        self.assertIn("Old notes", [s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide])
        self.assertIn("picture", r.stdout.lower())
        self.assertIn("2", r.stdout)

    def test_users_slide_template_is_used_and_its_sample_slides_dropped(self):
        own = Presentation()
        own.slides.add_slide(own.slide_layouts[0]).shapes.title.text = "TEMPLATE SAMPLE"
        own.save(self.tmp / "master.pptx")
        run(self.kit, "import_template.py", self.tmp / "master.pptx")
        (self.tmp / "deck.md").write_text("# Title\n\n---\n\n## Point\n- one\n")
        r = run(self.kit, "make_pptx.py", self.tmp / "deck.md", "-o", self.tmp / "deck.pptx")
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(self.tmp / "deck.pptx")
        self.assertEqual(len(prs.slides), 2)
        text = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
        self.assertNotIn("TEMPLATE SAMPLE", text)

    # ---------- self-contained check ----------
    def test_check_kit_passes_for_a_complete_kit_run_from_anywhere(self):
        self.templates()
        r = run(self.kit, "check_kit.py", cwd=self.tmp)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("ready to share", r.stdout.lower())

    def test_check_kit_flags_paths_outside_the_folder_and_broken_links(self):
        self.templates()
        (self.kit / "SKILL.md").write_text("---\nname: x\ndescription: y\n---\nUse /Users/someone/logo.png and [missing](assets/nope.png).\n")
        r = run(self.kit, "check_kit.py")
        self.assertEqual(r.returncode, 1)
        self.assertIn("/Users/someone", r.stdout)
        self.assertIn("assets/nope.png", r.stdout)

    def test_kit_folder_can_be_moved_and_still_works(self):
        self.templates()
        moved = self.tmp / "elsewhere" / "brandcreator-acme"
        shutil.move(str(self.kit), str(moved))
        (self.tmp / "memo.md").write_text("# Memo\n\nHi.\n")
        r = run(moved, "make_docx.py", self.tmp / "memo.md", "-o", self.tmp / "memo.docx")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(run(moved, "check_kit.py").returncode, 0)


if __name__ == "__main__":
    unittest.main()
